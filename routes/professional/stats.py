# noodle --uses-> services (export,pdf,xlsx)
# noodle --uses-> services (matching,coverage)
# noodle --queries-> database (sql,leads,professionals)
# noodle --renders-> templates (html,professional,lead_detail)
import csv
import io
import logging
import os
import re

from datetime import datetime
from io import StringIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
import pytz

logger = logging.getLogger(__name__)

from flask import Blueprint, flash, jsonify, redirect, render_template, request, Response, send_file, session, url_for
from fpdf import FPDF
from werkzeug.utils import secure_filename

import config
import models
import rate_limit
import utils
from decorators import login_required, professional_required
from i18n import t, get_language
from services.database import date_format_sql, now_sql
from services.export_helpers import pdf_safe, pdf_val, _style_header_row, _apply_data_border
from services.matching import lead_matches_coverage
from utils import allowed_file, convert_to_argentina_time


from routes.professional import professional_bp
from routes.professional.helpers import _get_approved_professional, _get_pro_geo_filter

def _query_leads_stats(conn, user_id, my_leads, month):
    """Helper: build stats data dict. conn must be open, caller owns close."""
    where_clauses = ['1=1']
    params = []

    if my_leads:
        conds, geo_params = _get_pro_geo_filter(conn, user_id)
        where_clauses.extend(conds)
        params.extend(geo_params)

    where_sql = ' AND '.join(where_clauses)

    cur_params = params + [month]
    cur_stats = conn.execute(f'''
        SELECT
            COUNT(*) as total,
            AVG(CAST(SUBSTR(budget, 1, INSTR(budget, ' ') - 1) AS REAL)) as avg_budget
        FROM leads WHERE {where_sql} AND {date_format_sql('timestamp', '%Y-%m')} = ?
    ''', cur_params).fetchone()

    cur_year, cur_mon = month.split('-')
    mon_int = int(cur_mon)
    prev_year = str(int(cur_year) - 1) if mon_int == 1 else cur_year
    prev_mon = f'12' if mon_int == 1 else f'{mon_int - 1:02d}'
    prev_month = f'{prev_year}-{prev_mon}'
    prev_params = params + [prev_month]
    prev_stats = conn.execute(f'''
        SELECT COUNT(*) as total FROM leads
        WHERE {where_sql} AND {date_format_sql('timestamp', '%Y-%m')} = ?
    ''', prev_params).fetchone()

    pt_params = params + [month]
    by_property_type = conn.execute(f'''
        SELECT property_type, COUNT(*) as count
        FROM leads WHERE {where_sql} AND {date_format_sql('timestamp', '%Y-%m')} = ?
        GROUP BY property_type ORDER BY count DESC
    ''', pt_params).fetchall()

    z_params = params + [month]
    by_zone = conn.execute(f'''
        SELECT zone, COUNT(*) as count
        FROM leads WHERE {where_sql} AND {date_format_sql('timestamp', '%Y-%m')} = ?
        GROUP BY zone ORDER BY count DESC LIMIT 10
    ''', z_params).fetchall()

    ot_params = params + [month]
    by_operation_type = conn.execute(f'''
        SELECT type, COUNT(*) as count
        FROM leads WHERE {where_sql} AND {date_format_sql('timestamp', '%Y-%m')} = ?
        GROUP BY type ORDER BY count DESC
    ''', ot_params).fetchall()

    trend = conn.execute(f'''
        SELECT {date_format_sql('timestamp', '%Y-%m')} as month, COUNT(*) as count
        FROM leads WHERE {where_sql}
        GROUP BY month ORDER BY month DESC LIMIT 6
    ''', params).fetchall()

    active_zone_count = conn.execute(f'''
        SELECT COUNT(DISTINCT zone) as cnt
        FROM leads WHERE {where_sql} AND {date_format_sql('timestamp', '%Y-%m')} = ? AND zone != ''
    ''', params + [month]).fetchone()[0]

    return {
        'total': cur_stats['total'] or 0,
        'avg_budget': round(cur_stats['avg_budget'] or 0, 0),
        'by_property_type': [{'label': r['property_type'], 'value': r['count']} for r in by_property_type],
        'by_zone': [{'label': r['zone'], 'value': r['count']} for r in by_zone],
        'by_operation_type': [{'label': r['type'], 'value': r['count']} for r in by_operation_type],
        'trend': [{'label': r['month'], 'value': r['count']} for r in reversed(trend)],
        'previous_month': {'total': prev_stats['total'] or 0},
        'active_zones': active_zone_count,
        'month': month,
    }


@professional_bp.route('/api/leads/stats')
@professional_required
def get_leads_stats():
    lang = get_language()
    conn = None
    try:
        conn = models.get_db_connection()
        user = conn.execute('SELECT username FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if not user:
            return jsonify({"success": False, "error": t('prof.access_denied', lang)}), 403

        professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
        if not professional or professional['status'] != 'approved':
            return jsonify({"success": False, "error": t('prof.account_pending', lang)}), 403

        my_leads = request.args.get('my_leads', '1').strip() == '1'
        month = request.args.get('month', '').strip()
        if not month or not re.match(r'^\d{4}-\d{2}$', month):
            month = datetime.now().strftime('%Y-%m')

        stats = _query_leads_stats(conn, session['user_id'], my_leads, month)
        return jsonify({"success": True, "stats": stats, "month": month})
    except Exception as e:
        logger.exception('Error en get_leads_stats')
        return jsonify({"success": False, "error": t('prof.internal_error', lang)}), 500
    finally:
        if conn:
            conn.close()


@professional_bp.route('/api/leads/stats/export')
@professional_required
def export_stats_csv():
    lang = get_language()
    conn = None
    try:
        conn = models.get_db_connection()
        user = conn.execute('SELECT username FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if not user:
            return t('prof.access_denied', lang), 403

        professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
        if not professional or professional['status'] != 'approved':
            return t('prof.account_pending', lang), 403

        my_leads = request.args.get('my_leads', '1').strip() == '1'
        month = request.args.get('month', '').strip()
        if not month or not re.match(r'^\d{4}-\d{2}$', month):
            month = datetime.now().strftime('%Y-%m')

        stats = _query_leads_stats(conn, session['user_id'], my_leads, month)

        def generate():
            data = StringIO()
            writer = csv.writer(data)

            # Summary section
            writer.writerow([t('prof.export_summary', lang)])
            writer.writerow([t('prof.export_month', lang), stats.get('month', month)])
            writer.writerow(['Total Leads', stats['total']])
            writer.writerow([t('prof.export_avg_budget', lang), f"${stats['avg_budget']:,.0f}"])
            writer.writerow([t('prof.export_vs_prev_month', lang), stats['previous_month']['total']])
            writer.writerow([t('prof.export_active_zones', lang), stats['active_zones']])
            writer.writerow([])

            # Property type section
            writer.writerow([t('prof.export_property_type_section', lang)])
            writer.writerow([t('prof.export_type', lang), t('prof.export_count', lang)])
            for pt in stats['by_property_type']:
                writer.writerow([pt['label'], pt['value']])
            writer.writerow([])

            # Zone section
            writer.writerow([t('prof.export_zones_section', lang)])
            writer.writerow([t('prof.export_zone', lang), t('prof.export_count', lang)])
            for z in stats['by_zone']:
                writer.writerow([z['label'], z['value']])
            writer.writerow([])

            # Operation type section
            writer.writerow([t('prof.export_operation_type_section', lang)])
            writer.writerow([t('prof.export_operation', lang), t('prof.export_count', lang)])
            for ot in stats['by_operation_type']:
                writer.writerow([ot['label'], ot['value']])
            writer.writerow([])

            # Trend section
            writer.writerow([t('prof.export_trend_section', lang)])
            writer.writerow([t('prof.export_month', lang), 'Leads'])
            for trend in stats['trend']:
                writer.writerow([trend['label'], trend['value']])

            yield data.getvalue()
            data.seek(0)
            data.truncate(0)

        filename = f'estadisticas_{month}_{datetime.now().strftime("%Y%m%d_%H%M")}.csv'
        return Response(
            generate(),
            mimetype='text/csv',
            headers={'Content-Disposition': f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        logger.exception('Error en export_stats_csv')
        return t('prof.internal_error', lang), 500
    finally:
        if conn:
            conn.close()


@professional_bp.route('/api/leads/stats/export/xlsx')
@professional_required
def export_stats_xlsx():
    lang = get_language()
    conn = None
    try:
        conn = models.get_db_connection()
        user = conn.execute('SELECT username FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if not user:
            return t('prof.access_denied', lang), 403

        professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
        if not professional or professional['status'] != 'approved':
            return t('prof.account_pending', lang), 403

        my_leads = request.args.get('my_leads', '1').strip() == '1'
        month = request.args.get('month', '').strip()
        if not month or not re.match(r'^\d{4}-\d{2}$', month):
            month = datetime.now().strftime('%Y-%m')

        stats = _query_leads_stats(conn, session['user_id'], my_leads, month)

        wb = openpyxl.Workbook()

        # Shared data font
        data_font = Font(name='Manrope', size=10, color='000410')
        label_font = Font(name='Manrope', size=10, bold=True, color='735A3A')
        title_font = Font(name='Manrope', size=10, bold=True, color='000410')

        # ─── Sheet 1: Resumen ───
        ws_resumen = wb.active
        ws_resumen.title = t('prof.export_summary_sheet', lang)
        resumen_data = [
            [t('prof.export_month', lang), stats.get('month', month)],
            ['Total Leads', stats['total']],
            [t('prof.export_avg_budget', lang), stats['avg_budget']],
            [t('prof.export_vs_prev_month', lang), stats['previous_month']['total']],
            [t('prof.export_active_zones', lang), stats['active_zones']],
        ]
        for col, h in enumerate([t('prof.export_metric', lang), t('prof.export_value', lang)], 1):
            cell = ws_resumen.cell(row=1, column=col, value=h)
        _style_header_row(ws_resumen, 2)
        ws_resumen.column_dimensions['A'].width = 28
        ws_resumen.column_dimensions['B'].width = 18

        for row_idx, row_data in enumerate(resumen_data, 2):
            for col, val in enumerate(row_data, 1):
                cell = ws_resumen.cell(row=row_idx, column=col, value=val)
                cell.font = data_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                if col == 1:
                    cell.font = label_font
            _apply_data_border(ws_resumen, row_idx, 2)

        # Separator + note rows
        note_row = len(resumen_data) + 3
        ws_resumen.cell(row=note_row, column=1, value=t('prof.export_generated_by', lang)).font = Font(
            name='Manrope', size=8, italic=True, color='A68A64'
        )

        # ─── Sheet 2: Tipo de Propiedad ───
        ws_pt = wb.create_sheet(t('prof.export_property_type_sheet', lang))
        for col, h in enumerate([t('prof.export_type', lang), t('prof.export_count', lang)], 1):
            ws_pt.cell(row=1, column=col, value=h)
        _style_header_row(ws_pt, 2)
        ws_pt.column_dimensions['A'].width = 32
        ws_pt.column_dimensions['B'].width = 14
        ws_pt.auto_filter.ref = f'A1:B{len(stats["by_property_type"]) + 1}'
        for row_idx, pt in enumerate(stats['by_property_type'], 2):
            c1 = ws_pt.cell(row=row_idx, column=1, value=pt['label'])
            c1.font = label_font
            c1.alignment = Alignment(horizontal='left', vertical='center')
            c2 = ws_pt.cell(row=row_idx, column=2, value=pt['value'])
            c2.font = data_font
            c2.alignment = Alignment(horizontal='center', vertical='center')
            _apply_data_border(ws_pt, row_idx, 2)

        # ─── Sheet 3: Zonas ───
        ws_z = wb.create_sheet(t('prof.export_zones_sheet', lang))
        for col, h in enumerate([t('prof.export_zone', lang), t('prof.export_count', lang)], 1):
            ws_z.cell(row=1, column=col, value=h)
        _style_header_row(ws_z, 2)
        ws_z.column_dimensions['A'].width = 34
        ws_z.column_dimensions['B'].width = 14
        ws_z.auto_filter.ref = f'A1:B{len(stats["by_zone"]) + 1}'
        for row_idx, z in enumerate(stats['by_zone'], 2):
            c1 = ws_z.cell(row=row_idx, column=1, value=z['label'])
            c1.font = label_font
            c1.alignment = Alignment(horizontal='left', vertical='center')
            c2 = ws_z.cell(row=row_idx, column=2, value=z['value'])
            c2.font = data_font
            c2.alignment = Alignment(horizontal='center', vertical='center')
            _apply_data_border(ws_z, row_idx, 2)

        # ─── Sheet 4: Tipo de Operación ───
        ws_ot = wb.create_sheet(t('prof.export_operation_type_sheet', lang))
        for col, h in enumerate([t('prof.export_operation', lang), t('prof.export_count', lang)], 1):
            ws_ot.cell(row=1, column=col, value=h)
        _style_header_row(ws_ot, 2)
        ws_ot.column_dimensions['A'].width = 34
        ws_ot.column_dimensions['B'].width = 14
        ws_ot.auto_filter.ref = f'A1:B{len(stats["by_operation_type"]) + 1}'
        for row_idx, ot in enumerate(stats['by_operation_type'], 2):
            c1 = ws_ot.cell(row=row_idx, column=1, value=ot['label'])
            c1.font = label_font
            c1.alignment = Alignment(horizontal='left', vertical='center')
            c2 = ws_ot.cell(row=row_idx, column=2, value=ot['value'])
            c2.font = data_font
            c2.alignment = Alignment(horizontal='center', vertical='center')
            _apply_data_border(ws_ot, row_idx, 2)

        # ─── Sheet 5: Tendencia Mensual ───
        ws_t = wb.create_sheet(t('prof.export_trend_sheet', lang))
        for col, h in enumerate([t('prof.export_month', lang), 'Leads'], 1):
            ws_t.cell(row=1, column=col, value=h)
        _style_header_row(ws_t, 2)
        ws_t.column_dimensions['A'].width = 28
        ws_t.column_dimensions['B'].width = 14
        ws_t.auto_filter.ref = f'A1:B{len(stats["trend"]) + 1}'
        for row_idx, trend in enumerate(stats['trend'], 2):
            c1 = ws_t.cell(row=row_idx, column=1, value=trend['label'])
            c1.font = label_font
            c1.alignment = Alignment(horizontal='left', vertical='center')
            c2 = ws_t.cell(row=row_idx, column=2, value=trend['value'])
            c2.font = data_font
            c2.alignment = Alignment(horizontal='center', vertical='center')
            _apply_data_border(ws_t, row_idx, 2)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f'estadisticas_{month}_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        logger.exception('Error en export_stats_xlsx')
        return t('prof.internal_error', lang), 500
    finally:
        if conn:
            conn.close()


@professional_bp.route('/api/leads/filter-options/invalidate', methods=['POST'])
@professional_required
def invalidate_filter_cache():
    from app_setup import filter_cache
    lang = get_language()
    filter_cache.invalidate()
    return jsonify({"success": True, "message": t('prof.cache_invalidated', lang)})


