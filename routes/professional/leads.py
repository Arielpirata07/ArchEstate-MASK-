# noodle --uses-> services (export,pdf,xlsx)
# noodle --uses-> services (matching,coverage)
# noodle --queries-> database (sql,leads,professionals)
# noodle --renders-> templates (html,professional,lead_detail)
import csv
import io
import logging
import os
import re

from datetime import datetime
from io import StringIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
import pytz

logger = logging.getLogger(__name__)

from flask import Blueprint, flash, jsonify, redirect, render_template, request, Response, send_file, session, url_for
from fpdf import FPDF
from werkzeug.utils import secure_filename

import config
import models
import rate_limit
import utils
from decorators import login_required, professional_required
from i18n import t, get_language
from services.database import date_format_sql, now_sql
from services.export_helpers import pdf_safe, pdf_val, _style_header_row, _apply_data_border
from services.matching import lead_matches_coverage, score_lead_for_coverage
from services.sla import SLA_RANK, _parse_ts, classify_sla_status
from utils import allowed_file, convert_to_argentina_time


from routes.professional import professional_bp
from routes.professional.helpers import _get_approved_professional, _get_pro_geo_filter

@professional_bp.route('/profesional')
@professional_required
def professional_view():
    conn = None
    try:
        conn = models.get_db_connection()
        user = conn.execute('SELECT username, doc_path FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if not user:
            return redirect(url_for('public.index'))

        professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
        if not professional or professional['status'] != 'approved':
            return render_template('professional.html', pending=True, doc_path=user['doc_path'])

        return render_template('professional.html', pending=False, doc_path=user['doc_path'])
    finally:
        if conn:
            conn.close()


@professional_bp.route('/profesional/lead/<int:lead_id>')
@professional_required
def lead_detail(lead_id):
    conn = None
    try:
        conn = models.get_db_connection()
        user = conn.execute('SELECT username FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if not user:
            return redirect(url_for('public.index'))

        professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
        if not professional or professional['status'] != 'approved':
            return render_template('professional.html', pending=True)

        lead = conn.execute('SELECT * FROM leads WHERE id = ?', (lead_id,)).fetchone()
    finally:
        if conn:
            conn.close()

    if not lead:
        return redirect(url_for('professional.professional_view'))

    lead_dict = dict(lead)
    lead_dict['timestamp'] = convert_to_argentina_time(lead_dict['timestamp'])
    phone_raw = lead_dict.get('phone') or ''
    lead_dict['phone_e164'] = utils.normalize_phone_to_e164(phone_raw)
    lead_dict['phone_is_mobile'] = bool(lead_dict['phone_e164'] and utils.is_whatsapp_capable(lead_dict['phone_e164']))
    from routes.professional.timeline import build_timeline
    timeline = build_timeline(lead_id, get_language())
    return render_template('lead_detail.html', lead=lead_dict, timeline=timeline)


@professional_bp.route('/api/leads')
@professional_required
def get_leads_api():
    lang = get_language()
    conn = None
    try:
        conn = models.get_db_connection()

        user = conn.execute('SELECT username FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if not user:
            return jsonify({"success": False, "error": t('prof.access_denied', lang)}), 403

        professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
        if not professional or professional['status'] != 'approved':
            return jsonify({"success": False, "error": t('prof.account_pending', lang)}), 403

        my_leads = request.args.get('my_leads', '1').strip() == '1'
        search = request.args.get('search', '').strip()
        type_filter = request.args.get('type', '').strip()
        prop_type = request.args.get('property_type', '').strip()
        zone_filter = request.args.get('zone', '').strip()
        min_budget = request.args.get('min_budget', '').strip()
        max_budget = request.args.get('max_budget', '').strip()
        budget_range = request.args.get('budget_range', '').strip()
        currency_filter = request.args.get('currency', '').strip()
        sort_by = request.args.get('sort', 'timestamp')
        sort_order = request.args.get('order', 'desc')
        time_range = request.args.get('time_range', '').strip()
        try:
            page = max(1, int(request.args.get('page', '1')))
        except ValueError:
            page = 1
        try:
            per_page = max(5, min(100, int(request.args.get('per_page', '25'))))
        except ValueError:
            per_page = 25
        offset = (page - 1) * per_page

        BUDGET_RANGES = {
            'hasta_200k': (0, 200000),
            '200k_500k': (200000, 500000),
            '500k_1m': (500000, 1000000),
            '1m_2m': (1000000, 2000000),
            'mas_2m': (2000000, None),
        }
        if budget_range and budget_range in BUDGET_RANGES:
            rng = BUDGET_RANGES[budget_range]
            min_budget = str(rng[0]) if rng[0] else ''
            max_budget = str(rng[1]) if rng[1] else ''

        query = 'SELECT * FROM leads WHERE 1=1'
        params = []

        if my_leads:
            conds, geo_params = _get_pro_geo_filter(conn, session['user_id'])
            for c in conds:
                query += f' AND {c}'
            params.extend(geo_params)

        if search:
            query += ' AND (zone LIKE ? OR email LIKE ? OR type LIKE ? OR budget LIKE ?)'
            search_param = f'%{search}%'
            params.extend([search_param, search_param, search_param, search_param])

        if type_filter:
            query += ' AND type = ?'
            params.append(type_filter)

        if prop_type:
            query += ' AND property_type = ?'
            params.append(prop_type)

        if zone_filter:
            query += ' AND zone LIKE ?'
            params.append(f'%{zone_filter}%')

        if min_budget:
            try:
                min_val = float(min_budget)
                query += " AND CAST(REPLACE(REPLACE(budget, '.', ''), ',', '') AS REAL) >= ?"
                params.append(min_val)
            except ValueError:
                logger.debug('min_budget inválido: %r', min_budget)

        if max_budget:
            try:
                max_val = float(max_budget)
                query += " AND CAST(REPLACE(REPLACE(budget, '.', ''), ',', '') AS REAL) <= ?"
                params.append(max_val)
            except ValueError:
                logger.debug('max_budget inválido: %r', max_budget)

        if currency_filter:
            query += ' AND currency = ?'
            params.append(currency_filter)

        valid_sort_fields = ['id', 'type', 'zone', 'budget', 'timestamp', 'email', 'match']
        if sort_by not in valid_sort_fields:
            sort_by = 'timestamp'

        order = 'DESC' if sort_order.lower() == 'desc' else 'ASC'
        if order not in ('ASC', 'DESC'):
            order = 'ASC'

        if time_range:
            try:
                tr = int(time_range)
                if tr > 0:
                    query += " AND timestamp >= datetime('now', ?)"
                    params.append(f'-{tr} days')
            except ValueError:
                logger.debug('time_range inválido: %r', time_range)

        count_query = f'SELECT COUNT(*) as total FROM ({query}) AS sub'
        total = conn.execute(count_query, params).fetchone()['total']

        kpi_row = conn.execute(f'''
            SELECT
                SUM(CASE WHEN lt.seen IS NULL OR lt.seen = 0 THEN 1 ELSE 0 END) as unseen,
                SUM(CASE WHEN lt.contacted = 1 THEN 1 ELSE 0 END) as contacted
            FROM ({query}) l
            LEFT JOIN lead_tracking lt ON lt.lead_id = l.id AND lt.professional_id = ?
        ''', params + [session['user_id']]).fetchone()
        kpi_total = total
        kpi_unseen = kpi_row['unseen'] or 0
        kpi_contacted = kpi_row['contacted'] or 0

        # sort=match: el score vive en Python (cobertura del profesional),
        # así que se traen los filtrados (tope de seguridad), se puntúa,
        # se ordena y se pagina en memoria. Default intacto: timestamp.
        match_sort = (sort_by == 'match')
        if match_sort:
            query += ' ORDER BY id DESC LIMIT 5000'
        else:
            query += f' ORDER BY {sort_by} {order}, id DESC LIMIT ? OFFSET ?'
            params.extend([per_page, offset])

        leads = conn.execute(query, params).fetchall()

        pro_row = conn.execute(
            'SELECT province FROM professionals WHERE user_id = ?', (session['user_id'],)
        ).fetchone()
        pro_province = pro_row['province'] if pro_row else ''
        coverage = models.get_professional_coverage(session['user_id'])

        # sort=match consciente del SLA (Fase 3): dentro del mismo score, los
        # leads sin primer contacto en riesgo (breached/at_risk) suben y los ya
        # contactados se hunden (no se re-empujan). El match sigue siendo la
        # clave primaria: el score nunca deja de ser descendente.
        contacted_ids = set()
        if match_sort:
            contacted_ids = {
                r['lead_id'] for r in conn.execute('SELECT lead_id FROM lead_first_contact').fetchall()
            }
            contacted_ids |= {
                r['lead_id'] for r in conn.execute(
                    'SELECT lead_id FROM lead_tracking WHERE professional_id = ? AND contacted = 1',
                    (session['user_id'],)).fetchall()
            }

        leads_list = []
        for lead in leads:
            lead_dict = dict(lead)
            lead_dict['timestamp'] = convert_to_argentina_time(lead_dict['timestamp'])
            phone_raw = lead_dict.get('phone') or ''
            phone_e164 = utils.normalize_phone_to_e164(phone_raw)
            lead_dict['phone_e164'] = phone_e164
            lead_dict['phone_is_mobile'] = bool(phone_e164 and utils.is_whatsapp_capable(phone_e164))
            lead_dict['matches_coverage'] = lead_matches_coverage(lead_dict, coverage, pro_province)
            if match_sort:
                lead_dict['match_score'] = score_lead_for_coverage(lead_dict, coverage, pro_province)
                lead_dict['sla_status'] = classify_sla_status(
                    _parse_ts(lead['timestamp']), lead_dict['id'] in contacted_ids)
            leads_list.append(lead_dict)

        if match_sort:
            leads_list.sort(
                key=lambda l: (l['match_score'], SLA_RANK[l['sla_status']], l['id']),
                reverse=True)
            leads_list = leads_list[offset:offset + per_page]

        professional_id = session['user_id']
        lead_ids = [lead['id'] for lead in leads_list]

        tracking_map = {}
        if lead_ids:
            placeholders = ','.join(['?'] * len(lead_ids))
            tracking_rows = conn.execute(
                f'SELECT lead_id, seen, contacted FROM lead_tracking WHERE professional_id = ? AND lead_id IN ({placeholders})',
                [professional_id] + lead_ids
            ).fetchall()
            for row in tracking_rows:
                tracking_map[row['lead_id']] = {
                    'seen': bool(row['seen']),
                    'contacted': bool(row['contacted'])
                }

        for lead in leads_list:
            tracking = tracking_map.get(lead['id'], {'seen': False, 'contacted': False})
            lead['tracking'] = tracking

        total_pages = max(1, (total + per_page - 1) // per_page)

        return jsonify({
            "success": True,
            "leads": leads_list,
            "total": total,
            "page": page,
            "per_page": per_page,
            "total_pages": total_pages,
            "kpi_total": kpi_total,
            "kpi_unseen": kpi_unseen,
            "kpi_contacted": kpi_contacted
        })
    except Exception as e:
        logger.exception('Error en get_leads_api')
        return jsonify({"success": False, "error": t('prof.internal_error', lang)}), 500
    finally:
        if conn:
            conn.close()


@professional_bp.route('/api/leads/filter-options')
@professional_required
def get_leads_filter_options():
    lang = get_language()
    conn = None
    try:
        conn = models.get_db_connection()
        user = conn.execute('SELECT username FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if user:
            professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
            if not professional or professional['status'] != 'approved':
                return jsonify({"success": False, "error": t('prof.approval_required', lang)}), 403
    finally:
        if conn:
            conn.close()

    from app_setup import filter_cache
    cached = filter_cache.get('filter_options')
    if cached:
        return jsonify(cached)

    conn = None
    try:
        conn = models.get_db_connection()
        types = [r[0] for r in conn.execute('SELECT DISTINCT type FROM leads WHERE type IS NOT NULL ORDER BY type').fetchall()]
        prop_types = [r[0] for r in conn.execute('SELECT DISTINCT property_type FROM leads WHERE property_type IS NOT NULL ORDER BY property_type').fetchall()]
        currencies = [r[0] for r in conn.execute('SELECT DISTINCT currency FROM leads WHERE currency IS NOT NULL ORDER BY currency').fetchall()]
        zones = [r[0] for r in conn.execute('SELECT DISTINCT zone FROM leads WHERE zone IS NOT NULL ORDER BY zone').fetchall()]

        result = {
            'types': types,
            'property_types': prop_types,
            'currencies': currencies,
            'zones': zones,
        }
        filter_cache.set('filter_options', result)
        return jsonify(result)
    finally:
        if conn:
            conn.close()


@professional_bp.route('/api/leads/panorama')
@professional_required
def get_leads_panorama():
    lang = get_language()
    conn = None
    try:
        conn = models.get_db_connection()
        _, error = _get_approved_professional(conn, session['user_id'], lang)
        if error:
            return error
        total = conn.execute('SELECT COUNT(*) AS n FROM leads').fetchone()['n']

        def top(column):
            if column not in ('zone', 'province', 'country', 'type', 'property_type'):
                raise ValueError(f'Columna no permitida en top(): {column!r}')
            return [dict(r) for r in conn.execute(
                f"SELECT {column} AS label, COUNT(*) AS value FROM leads "
                f"WHERE {column} IS NOT NULL AND {column} != '' "
                f'GROUP BY {column} ORDER BY value DESC LIMIT 8'
            ).fetchall()]

        my_total = None
        try:
            conds, geo_params = _get_pro_geo_filter(conn, session['user_id'])
            where = ' AND '.join(conds) if conds else '1=1'
            my_total = conn.execute(
                f'SELECT COUNT(*) AS n FROM leads WHERE {where}', geo_params
            ).fetchone()['n']
        except Exception:
            logger.debug('my_total falló, se omite')
        return jsonify({
            'success': True,
            'panorama': {
                'total': total,
                'my_total': my_total,
                'by_zone': top('zone'),
                'by_province': top('province'),
                'by_country': top('country'),
            }
        })
    except Exception:
        logger.exception('Error en get_leads_panorama')
        return jsonify({'success': False, 'error': t('prof.internal_error_short', lang)}), 500
    finally:
        if conn:
            conn.close()


@professional_bp.route('/api/professional/my-org')
@professional_required
def get_my_org():
    lang = get_language()
    conn = None
    try:
        conn = models.get_db_connection()
        _, error = _get_approved_professional(conn, session['user_id'], lang)
        if error:
            return error
        org = models.get_my_org(session['user_id'])
        if not org:
            return jsonify({'success': True, 'org': None, 'message': t('org.no_org', lang)})
        members = models.get_org_members(org['id'])
        return jsonify({'success': True, 'org': org, 'members': members})
    except Exception:
        logger.exception('Error en get_my_org')
        return jsonify({'success': False, 'error': t('prof.internal_error_short', lang)}), 500
    finally:
        if conn:
            conn.close()


@professional_bp.route('/api/professional/org-leads')
@professional_required
def get_org_leads():
    lang = get_language()
    conn = None
    try:
        conn = models.get_db_connection()
        _, error = _get_approved_professional(conn, session['user_id'], lang)
        if error:
            return error
        org = models.get_my_org(session['user_id'])
        if not org:
            return jsonify({'success': True, 'leads': [], 'total': 0, 'message': t('org.no_org', lang)})
        rows = conn.execute(
            'SELECT * FROM leads WHERE org_id = ? ORDER BY timestamp DESC LIMIT 100',
            (org['id'],)
        ).fetchall()
        return jsonify({'success': True, 'org_id': org['id'], 'leads': [dict(r) for r in rows], 'total': len(rows)})
    except Exception:
        logger.exception('Error en get_org_leads')
        return jsonify({'success': False, 'error': t('prof.internal_error_short', lang)}), 500
    finally:
        if conn:
            conn.close()


