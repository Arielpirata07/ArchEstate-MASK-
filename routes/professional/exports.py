# noodle --uses-> services (export,pdf,xlsx)
# noodle --uses-> services (matching,coverage)
# noodle --queries-> database (sql,leads,professionals)
# noodle --renders-> templates (html,professional,lead_detail)
import csv
import io
import logging
import os
import re

from datetime import datetime
from io import StringIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
import pytz

logger = logging.getLogger(__name__)

from flask import Blueprint, flash, jsonify, redirect, render_template, request, Response, send_file, session, url_for
from fpdf import FPDF
from werkzeug.utils import secure_filename

import config
import models
import rate_limit
import utils
from decorators import login_required, professional_required
from i18n import t, get_language
from services.database import date_format_sql, now_sql
from services.export_helpers import pdf_safe, pdf_val, _style_header_row, _apply_data_border
from services.matching import lead_matches_coverage
from utils import allowed_file, convert_to_argentina_time


from routes.professional import professional_bp
from routes.professional.helpers import _get_approved_professional, _get_pro_geo_filter

@professional_bp.route('/api/leads/export')
@professional_required
@rate_limit.check_rate_limit(limit=20, window=60)
def export_leads_csv():
    conn = None
    leads = []
    try:
        conn = models.get_db_connection()
        lang = get_language()
        user = conn.execute('SELECT username FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if not user:
            return t('prof.access_denied', lang), 403

        professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
        if not professional or professional['status'] != 'approved':
            return t('prof.account_pending', lang), 403

        conds, geo_params = _get_pro_geo_filter(conn, session['user_id'])
        query = 'SELECT id, type, zone, budget, currency, timestamp FROM leads WHERE 1=1'
        params = []
        for c in conds:
            query += f' AND {c}'
        params.extend(geo_params)
        query += ' ORDER BY timestamp DESC'
        leads = conn.execute(query, params).fetchall()
    finally:
        if conn:
            conn.close()

    def generate():
        data = StringIO()
        writer = csv.writer(data)
        writer.writerow(t('prof.export_csv_headers', lang).split(','))
        yield data.getvalue()
        data.seek(0)
        data.truncate(0)

        for lead in leads:
            timestamp_argentina = convert_to_argentina_time(lead['timestamp'])
            writer.writerow([lead['id'], lead['type'], lead['zone'], lead['budget'], lead['currency'], timestamp_argentina])
            yield data.getvalue()
            data.seek(0)
            data.truncate(0)

    filename = f"leads_archestate_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    return Response(
        generate(),
        mimetype='text/csv',
        headers={
            'Content-Disposition': f'attachment; filename="{filename}"'
        }
    )


@professional_bp.route('/api/leads/export/xlsx')
@professional_required
@rate_limit.check_rate_limit(limit=20, window=60)
def export_leads_xlsx():
    conn = None
    leads = []
    try:
        conn = models.get_db_connection()
        lang = get_language()
        user = conn.execute('SELECT username FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if not user:
            return t('prof.access_denied', lang), 403

        professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
        if not professional or professional['status'] != 'approved':
            return t('prof.account_pending', lang), 403

        conds, geo_params = _get_pro_geo_filter(conn, session['user_id'])
        query = 'SELECT id, type, zone, budget, currency, timestamp FROM leads WHERE 1=1'
        params = []
        for c in conds:
            query += f' AND {c}'
        params.extend(geo_params)
        query += ' ORDER BY timestamp DESC'
        leads = conn.execute(query, params).fetchall()
    finally:
        if conn:
            conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = t('prof.export_leads_sheet', lang)

    headers = t('prof.export_leads_headers', lang).split(',')
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)
    _style_header_row(ws, len(headers))

    col_widths = [10, 30, 28, 18, 12, 24]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    data_font = Font(name='Manrope', size=10, color='000410')
    data_align = Alignment(horizontal='center', vertical='center')
    date_font = Font(name='Manrope', size=10, color='A68A64')

    for row_num, lead in enumerate(leads, 2):
        timestamp_argentina = convert_to_argentina_time(lead['timestamp'])
        vals = [lead['id'], lead['type'], lead['zone'], lead['budget'], lead['currency'], timestamp_argentina]
        for col_num, val in enumerate(vals, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = data_font
            cell.alignment = data_align
            if col_num == len(vals):
                cell.font = date_font
        _apply_data_border(ws, row_num, len(headers))

    ws.auto_filter.ref = f'A1:F{len(leads) + 1}'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"leads_archestate_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


@professional_bp.route('/api/lead/<int:lead_id>/download')
@professional_required
def download_lead_pdf(lead_id):
    lang = get_language()
    conn = None
    try:
        conn = models.get_db_connection()
        user = conn.execute('SELECT username FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        if not user:
            return t('prof.access_denied', lang), 403

        professional = conn.execute('SELECT status FROM professionals WHERE name = ?', (user['username'],)).fetchone()
        if not professional or professional['status'] != 'approved':
            return t('prof.account_pending', lang), 403

        lead = conn.execute('SELECT * FROM leads WHERE id = ?', (lead_id,)).fetchone()
    finally:
        if conn:
            conn.close()

    if not lead:
        return jsonify({"success": False, "error": t('prof.lead_not_found', lang)}), 404

    lead = dict(lead)

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    midnight = (0, 4, 16)
    gold = (115, 90, 58)

    pdf.set_font('Times', 'BI', 20)
    pdf.set_text_color(*midnight)
    pdf.cell(0, 15, t('prof.pdf_title', lang), ln=True, align='C')
    pdf.ln(5)

    pdf.set_font('Helvetica', '', 10)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 8, t('prof.pdf_subtitle', lang, lead_id=lead['id']), ln=True, align='C')
    pdf.ln(10)

    def section_header(title):
        pdf.set_fill_color(*gold)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font('Helvetica', 'B', 12)
        pdf.cell(0, 8, title.upper(), ln=True, fill=True)
        pdf.set_text_color(*midnight)
        pdf.set_font('Helvetica', '', 10)
        pdf.ln(2)

    section_header(t('prof.pdf_operation_type', lang))
    pdf.cell(0, 6, pdf_val(lead['type']), ln=True)

    prop_type = pdf_safe(lead.get('property_type', '')).lower()
    prop_type_label = {'departamento': 'Departamento', 'casa': 'Casa', 'duplex': 'Duplex',
                       'penthouse': 'Penthouse', 'local_comercial': 'Local Comercial'}.get(prop_type, prop_type)
    section_header(t('prof.pdf_property_type', lang))
    pdf.cell(0, 6, pdf_val(prop_type_label, t('prof.not_specified', lang)), ln=True)

    section_header(t('prof.pdf_geographic_zone', lang))
    pdf.cell(0, 6, f"{t('prof.pdf_zone_prefix', lang)}{pdf_val(lead['zone'])}", ln=True)
    pdf.cell(0, 6, f"{t('prof.pdf_province_prefix', lang)}{pdf_val(lead.get('province'), t('prof.not_specified_fem', lang))}", ln=True)

    section_header(t('prof.pdf_budget', lang))
    budget_symbol = 'USD' if lead['currency'] == 'USD' else 'EUR' if lead['currency'] == 'EUR' else '$'
    pdf.cell(0, 6, f"{budget_symbol} {pdf_val(lead['budget'])}", ln=True)

    section_header(t('prof.pdf_architectural_style', lang))
    pdf.cell(0, 6, pdf_val(lead.get('architectural_style'), t('prof.not_specified', lang)), ln=True)

    section_header(t('prof.pdf_direct_contact', lang))
    pdf.cell(0, 6, f"Email: {pdf_val(lead['email'])}", ln=True)
    pdf.cell(0, 6, f"Telefono: {pdf_val(lead['phone'])}", ln=True)

    section_header(t('prof.pdf_registered', lang))
    pdf.cell(0, 6, pdf_val(convert_to_argentina_time(lead['timestamp'])), ln=True)
    pdf.ln(5)

    section_header(t('prof.pdf_tech_specs', lang))

    pdf.set_font('Helvetica', 'B', 10)
    pdf.cell(60, 8, t('prof.pdf_rooms', lang), border=1)
    pdf.cell(0, 8, pdf_val(lead.get('ambientes')), ln=True, border=1)

    pdf.cell(60, 8, t('prof.pdf_bedrooms', lang), border=1)
    pdf.cell(0, 8, pdf_val(lead['bedrooms']), ln=True, border=1)

    pdf.cell(60, 8, t('prof.pdf_bathrooms', lang), border=1)
    pdf.cell(0, 8, pdf_val(lead['bathrooms']), ln=True, border=1)

    pdf.cell(60, 8, t('prof.pdf_parking', lang), border=1)
    pdf.cell(0, 8, pdf_val(lead.get('parking'), t('prof.not_specified', lang)), ln=True, border=1)

    pdf.cell(60, 8, t('prof.pdf_orientation', lang), border=1)
    pdf.cell(0, 8, pdf_val(lead.get('orientation'), t('prof.not_specified_fem', lang)), ln=True, border=1)

    pdf.ln(2)

    pdf.cell(60, 8, t('prof.pdf_usable_m2', lang), border=1)
    pdf.cell(0, 8, f"{pdf_val(lead.get('usable_m2'))} m2" if lead.get('usable_m2') else '-', ln=True, border=1)

    pdf.cell(60, 8, t('prof.pdf_total_area', lang), border=1)
    pdf.cell(0, 8, f"{pdf_val(lead.get('total_area'))} m2" if lead.get('total_area') else '-', ln=True, border=1)

    pdf.cell(60, 8, t('prof.pdf_land_area', lang), border=1)
    pdf.cell(0, 8, f"{pdf_val(lead.get('land_area'))} m2" if lead.get('land_area') else '-', ln=True, border=1)

    pdf.cell(60, 8, t('prof.pdf_built_area', lang), border=1)
    pdf.cell(0, 8, f"{pdf_val(lead.get('built_area'))} m2" if lead.get('built_area') else '-', ln=True, border=1)

    pdf.ln(5)

    section_header(t('prof.pdf_condition_age', lang))
    pdf.set_font('Helvetica', 'B', 10)
    pdf.cell(60, 8, t('prof.pdf_condition', lang), border=1)
    pdf.cell(0, 8, pdf_val(lead.get('property_condition'), t('prof.not_specified', lang)), ln=True, border=1)

    pdf.cell(60, 8, t('prof.pdf_age', lang), border=1)
    pdf.cell(0, 8, pdf_val(lead.get('property_age'), t('prof.not_specified_fem', lang)), ln=True, border=1)

    pdf.ln(5)

    section_header(t('prof.pdf_amenities', lang))
    amenities = lead.get('amenities', '')
    if amenities and str(amenities).strip():
        for amenity in pdf_safe(amenities).split(','):
            stripped = amenity.strip()
            if stripped:
                pdf.cell(0, 6, f"- {stripped}", ln=True)
    else:
        pdf.cell(0, 6, t('prof.not_specified_fem_pl', lang), ln=True)

    section_header(t('prof.pdf_property_details', lang))
    if prop_type == 'departamento':
        pdf.cell(0, 6, f"{t('prof.pdf_floor_block', lang)}{pdf_val(lead.get('floor_block'), t('prof.not_specified', lang))}", ln=True)
    pdf.cell(0, 6, f"{t('prof.pdf_pool', lang)}{pdf_val(lead.get('pool'), t('prof.not_specified_fem', lang))}", ln=True)
    pdf.cell(0, 6, f"{t('prof.pdf_elevator', lang)}{pdf_val(lead.get('elevator'), t('prof.not_specified', lang))}", ln=True)

    pdf_output = pdf.output(dest='S')
    if isinstance(pdf_output, str):
        pdf_output = pdf_output.encode('latin-1')

    buffer = io.BytesIO(pdf_output)
    buffer.seek(0)

    filename = f"lead_{lead['id']}.pdf"

    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )


