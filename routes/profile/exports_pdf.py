# noodle --uses-> services (export,pdf,xlsx)
# noodle --uses-> services (notifications,email)
# noodle --queries-> database (sql,users,leads,professionals,coverage)
# noodle --renders-> templates (html,profile,edit_lead,notifications)
import os
import json
import logging
import time
from datetime import datetime
from io import BytesIO

from flask import Blueprint, render_template, request, session, jsonify, redirect, url_for, current_app, send_file
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

import config
import models
import decorators
import validators
import utils
from utils import parse_budget
import rate_limit
from services.export_helpers import pdf_safe, pdf_val, _style_header_row, _apply_data_border
from i18n import t, get_language
from services.notifications import notify_admins


logger = logging.getLogger(__name__)


profile_bp = Blueprint('profile', __name__, url_prefix='')


ALLOWED_LEAD_EDIT_FIELDS = [
    'zone', 'province', 'budget', 'currency',
    'floor_block', 'usable_m2', 'elevator',
    'land_area', 'built_area', 'pool',
    'architectural_style', 'bedrooms', 'bathrooms',
    'total_area', 'amenities', 'ambientes',
    'parking', 'orientation', 'property_condition',
    'property_age', 'community_pool', 'additional_features'
]


from routes.profile import profile_bp
from routes.profile.exports import _build_export_stats, _query_monthly_contacts

@profile_bp.route('/api/profile/professional/export/pdf')
@decorators.professional_required
@rate_limit.check_rate_limit(limit=20, window=60)
def profile_export_history_pdf():
    try:
        lang = get_language()
        user_id = session['user_id']
        leads = _query_monthly_contacts(user_id)
        stats = _build_export_stats(leads)

        # Fetch professional name
        conn = models.get_db_connection()
        pro = conn.execute(
            'SELECT name FROM professionals WHERE user_id = ?', (user_id,)
        ).fetchone()
        prof_name = pdf_safe(pro['name']) if pro else 'Profesional'
        conn.close()

        midnight = (0, 4, 16)
        gold = (115, 90, 58)
        gold_light = (166, 138, 100)
        gray_bg = (245, 243, 240)
        white = (255, 255, 255)

        class ReportPDF(FPDF):
            def header(self):
                if self.page_no() == 1:
                    return
                self.set_fill_color(*midnight)
                self.rect(0, 0, 210, 10, 'F')
                self.set_font('Helvetica', 'B', 7)
                self.set_text_color(*gold_light)
                self.set_y(2)
                self.cell(0, 6, t('profile.pdf_report_title', lang), ln=True, align='C')
                self.ln(4)

            def footer(self):
                self.set_y(-12)
                self.set_font('Helvetica', 'I', 7)
                self.set_text_color(150, 150, 150)
                self.cell(0, 8, f'{t("profile.pdf_page", lang)} {self.page_no()}/{{nb}}  |  {datetime.now().strftime("%d/%m/%Y %H:%M")}', ln=True, align='C')

        pdf = ReportPDF()
        pdf.alias_nb_pages()
        pdf.set_auto_page_break(auto=True, margin=18)

        # ─── PAGE 1 — Brand Header ───
        pdf.add_page()
        pdf.set_fill_color(*midnight)
        pdf.rect(0, 0, 210, 28, 'F')

        pdf.set_y(4)
        pdf.set_font('Times', 'BI', 18)
        pdf.set_text_color(*gold_light)
        pdf.cell(0, 10, 'ArchEstate', ln=True, align='C')
        pdf.set_font('Helvetica', '', 7)
        pdf.set_text_color(166, 138, 100)
        pdf.cell(0, 5, 'The Private Ledger', ln=True, align='C')

        pdf.set_y(24)
        pdf.set_draw_color(*gold)
        pdf.set_line_width(0.5)
        pdf.line(20, 28, 190, 28)

        # ─── Professional info ───
        pdf.set_y(32)
        pdf.set_font('Helvetica', 'B', 10)
        pdf.set_text_color(*midnight)
        pdf.cell(0, 6, pdf_safe(prof_name), ln=True, align='L')
        pdf.set_font('Helvetica', '', 8)
        pdf.set_text_color(100, 100, 100)
        pdf.cell(0, 5, t('profile.pdf_subtitle', lang), ln=True, align='L')
        pdf.ln(6)

        # ─── Summary cards (2x2) ───
        card_w = 82
        card_h = 18
        gap = 6
        left_margin = 14
        pdf.set_font('Helvetica', '', 9)

        cards = [
            (t('profile.pdf_total_leads', lang), str(stats['total']), f"{stats['zone_count']} {t('profile.pdf_active_zones', lang)}"),
            (t('profile.pdf_total_budget', lang), f"${stats['total_budget']:,.2f}", f"{t('profile.pdf_average', lang)} ${stats['avg_budget']:,.2f}"),
        ]
        pdf.set_x(left_margin)
        for title, value, sub in cards:
            x0 = pdf.get_x()
            y0 = pdf.get_y()
            pdf.set_fill_color(*gold)
            pdf.rect(x0, y0, 4, card_h, 'F')
            pdf.set_fill_color(*gray_bg)
            pdf.rect(x0 + 4, y0, card_w - 4, card_h, 'F')
            pdf.set_text_color(*midnight)
            pdf.set_xy(x0 + 8, y0 + 2)
            pdf.set_font('Helvetica', '', 7)
            pdf.cell(card_w - 12, 4, pdf_safe(title))
            pdf.set_xy(x0 + 8, y0 + 6)
            pdf.set_font('Helvetica', 'B', 13)
            pdf.cell(card_w - 12, 7, pdf_safe(value))
            pdf.set_xy(x0 + 8, y0 + 13)
            pdf.set_font('Helvetica', 'I', 7)
            pdf.set_text_color(100, 100, 100)
            pdf.cell(card_w - 12, 4, pdf_safe(sub))
            pdf.set_x(x0 + card_w + gap)

        pdf.ln(card_h + 4)
        pdf.set_text_color(*midnight)

        pdf.set_x(left_margin)
        cards2 = [
            (t('profile.pdf_avg_per_lead', lang), f"${stats['avg_budget']:,.2f}", t('profile.pdf_estimated_avg', lang)),
            (t('profile.pdf_active_currencies', lang), str(len(stats['currencies'])), f"{', '.join(sorted(stats['currencies'].keys()))}"),
        ]
        for title, value, sub in cards2:
            x0 = pdf.get_x()
            y0 = pdf.get_y()
            pdf.set_fill_color(*gold)
            pdf.rect(x0, y0, 4, card_h, 'F')
            pdf.set_fill_color(*gray_bg)
            pdf.rect(x0 + 4, y0, card_w - 4, card_h, 'F')
            pdf.set_text_color(*midnight)
            pdf.set_xy(x0 + 8, y0 + 2)
            pdf.set_font('Helvetica', '', 7)
            pdf.cell(card_w - 12, 4, pdf_safe(title))
            pdf.set_xy(x0 + 8, y0 + 6)
            pdf.set_font('Helvetica', 'B', 13)
            pdf.cell(card_w - 12, 7, pdf_safe(value))
            pdf.set_xy(x0 + 8, y0 + 13)
            pdf.set_font('Helvetica', 'I', 7)
            pdf.set_text_color(100, 100, 100)
            pdf.cell(card_w - 12, 4, pdf_safe(sub))
            pdf.set_x(x0 + card_w + gap)

        pdf.ln(card_h + 8)

        # ─── Currency breakdown ───
        def section_header(title):
            pdf.set_fill_color(*gold)
            pdf.set_text_color(*white)
            pdf.set_font('Helvetica', 'B', 10)
            pdf.cell(0, 7, pdf_safe('  ' + title.upper()), ln=True, fill=True)
            pdf.set_text_color(*midnight)
            pdf.set_font('Helvetica', '', 9)
            pdf.ln(2)

        section_header(t('profile.pdf_by_currency', lang))
        pdf.set_font('Helvetica', 'B', 8)
        pdf.set_fill_color(*gray_bg)
        cw = [30, 30, 50, 50, 30]
        ch = [t('profile.pdf_currency', lang), t('profile.pdf_count', lang), t('profile.pdf_total_budget_label', lang), t('profile.pdf_avg_budget_label', lang), t('profile.pdf_pct', lang)]
        for i, (h, w) in enumerate(zip(ch, cw)):
            pdf.cell(w, 6, pdf_safe(h), border=1, align='C', fill=True)
        pdf.ln()
        pdf.set_font('Helvetica', '', 8)
        for currency, count in sorted(stats['currencies'].items()):
            bc = stats['budget_by_currency'].get(currency, {'count': 0, 'total': 0.0})
            total_c = bc['total']
            avg_c = round(total_c / bc['count'], 2) if bc['count'] else 0
            pct = round(count / stats['total'] * 100, 1) if stats['total'] else 0
            vals = [currency, str(count), f"${total_c:,.2f}", f"${avg_c:,.2f}", f"{pct}%"]
            for v, w in zip(vals, cw):
                pdf.cell(w, 5, pdf_safe(v), border=1, align='C')
            pdf.ln()
        pdf.ln(5)

        # ─── Property type breakdown ───
        section_header(t('profile.pdf_by_property_type', lang))
        pdf.set_font('Helvetica', 'B', 8)
        pdf.set_fill_color(*gray_bg)
        cw2 = [80, 40, 40, 30]
        ch2 = [t('profile.pdf_type', lang), t('profile.pdf_count', lang), t('profile.pdf_pct_total', lang), t('profile.pdf_avg_budget', lang)]
        for h, w in zip(ch2, cw2):
            pdf.cell(w, 6, pdf_safe(h), border=1, align='C', fill=True)
        pdf.ln()
        pdf.set_font('Helvetica', '', 8)

        # Compute per-type budgets
        type_budgets = {}
        for lead in leads:
            type_val = lead.get('type') or 'Otro'
            b = parse_budget(lead.get('budget'))
            tb = type_budgets.setdefault(type_val, {'count': 0, 'total': 0.0})
            tb['count'] += 1
            tb['total'] += b
        for type_val, count in sorted(stats['types'].items(), key=lambda x: x[1], reverse=True):
            tb = type_budgets.get(type_val, {'count': 0, 'total': 0.0})
            avg_t = round(tb['total'] / tb['count'], 2) if tb['count'] else 0
            pct = round(count / stats['total'] * 100, 1) if stats['total'] else 0
            vals = [pdf_safe(type_val)[:35], str(count), f"{pct}%", f"${avg_t:,.2f}"]
            for v, w in zip(vals, cw2):
                pdf.cell(w, 5, pdf_safe(v), border=1, align='C')
            pdf.ln()
        pdf.ln(5)

        # ─── Zone breakdown ───
        section_header(t('profile.pdf_top_zones', lang))
        pdf.set_font('Helvetica', 'B', 8)
        pdf.set_fill_color(*gray_bg)
        cw3 = [90, 40, 40, 22]
        ch3 = [t('profile.pdf_zone', lang), t('profile.pdf_count', lang), t('profile.pdf_pct_total', lang), t('profile.pdf_bar', lang)]
        for h, w in zip(ch3, cw3):
            pdf.cell(w, 6, pdf_safe(h), border=1, align='C', fill=True)
        pdf.ln()
        pdf.set_font('Helvetica', '', 8)
        max_count = max((c for _, c in stats['zones']), default=1)
        for zone, count in stats['zones'][:10]:
            pct = round(count / stats['total'] * 100, 1) if stats['total'] else 0
            vals = [pdf_safe(zone)[:40], str(count), f"{pct}%"]
            for v, w in zip(vals, cw3[:3]):
                pdf.cell(w, 5, pdf_safe(v), border=1, align='C')
            # Mini bar
            bar_ratio = count / max_count
            bar_w = int(bar_ratio * 18)
            x0 = pdf.get_x()
            y0 = pdf.get_y()
            pdf.set_fill_color(*gold_light)
            pdf.rect(x0, y0 + 1, bar_w, 3, 'F')
            pdf.cell(cw3[3], 5, '', border=1)
            pdf.ln()
        pdf.ln(5)

        # ─── Leads detail table ───
        section_header(t('profile.pdf_lead_detail', lang))
        pdf.set_font('Helvetica', 'B', 7)
        pdf.set_fill_color(*midnight)
        pdf.set_text_color(*white)
        cols_l = [8, 8, 32, 26, 24, 22, 28, 22]
        headers_l = ['#', 'ID', t('profile.pdf_type', lang), t('profile.pdf_zone', lang), t('profile.pdf_budget', lang), t('profile.pdf_currency', lang), t('profile.pdf_province', lang), t('profile.pdf_contacted_label', lang)]
        for h, w in zip(headers_l, cols_l):
            pdf.cell(w, 6, pdf_safe(h), border=1, align='C', fill=True)
        pdf.ln()
        pdf.set_text_color(*midnight)
        pdf.set_font('Helvetica', '', 7)
        for idx, lead in enumerate(leads[:30], 1):
            # contacted_at es UTC en storage: convertir antes de cortar la
            # fecha, si no el corte [0:10] a veces daba el día UTC, no el local.
            contact_date = utils.convert_to_argentina_time(lead['contacted_at'])[:10] if lead.get('contacted_at') else ''
            budget = pdf_val(lead['budget'])
            vals = [
                str(idx), str(lead['id']), pdf_safe(lead['type'])[:28],
                pdf_safe(lead['zone'])[:22],
                budget[:20], pdf_safe(lead.get('currency', ''))[:8],
                pdf_safe(lead.get('province', ''))[:24],
                contact_date[:20],
            ]
            if idx % 2 == 0:
                pdf.set_fill_color(*gray_bg)
                fill = True
            else:
                fill = False
            for v, w in zip(vals, cols_l):
                pdf.cell(w, 5, pdf_safe(v), border=1, align='C', fill=fill)
            pdf.ln()
            if pdf.get_y() > 258:
                pdf.add_page()

        pdf.ln(5)
        pdf.set_font('Helvetica', 'I', 7)
        pdf.set_text_color(*gold_light)
        gen_date = datetime.now().strftime('%d/%m/%Y %H:%M')
        pdf.cell(0, 5, f'{t("profile.pdf_generated_by", lang)} | {gen_date}', ln=True, align='C')

        pdf_output = pdf.output(dest='S')
        if isinstance(pdf_output, str):
            pdf_output = pdf_output.encode('latin-1')

        buffer = BytesIO(pdf_output)
        buffer.seek(0)

        now_str = datetime.now().strftime('%Y%m%d_%H%M')
        return send_file(
            buffer,
            as_attachment=True,
            download_name=f'reporte_rendimiento_{now_str}.pdf',
            mimetype='application/pdf',
        )
    except Exception:
        logger.exception('Error exporting PDF report')
        lang = get_language()
        return jsonify({'error': t('profile.pdf_export_error', lang)}), 500
