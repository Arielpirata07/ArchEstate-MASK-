# noodle --uses-> services (export,pdf,xlsx)
# noodle --uses-> services (notifications,email)
# noodle --queries-> database (sql,users,leads,professionals,coverage)
# noodle --renders-> templates (html,profile,edit_lead,notifications)
import os
import json
import logging
import time
from datetime import datetime
from io import BytesIO

from flask import Blueprint, render_template, request, session, jsonify, redirect, url_for, current_app, send_file
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

import config
import models
import decorators
import validators
import utils
from utils import parse_budget
import rate_limit
from services.export_helpers import pdf_safe, pdf_val, _style_header_row, _apply_data_border
from i18n import t, get_language
from services.notifications import notify_admins


logger = logging.getLogger(__name__)


profile_bp = Blueprint('profile', __name__, url_prefix='')


ALLOWED_LEAD_EDIT_FIELDS = [
    'zone', 'province', 'budget', 'currency',
    'floor_block', 'usable_m2', 'elevator',
    'land_area', 'built_area', 'pool',
    'architectural_style', 'bedrooms', 'bathrooms',
    'total_area', 'amenities', 'ambientes',
    'parking', 'orientation', 'property_condition',
    'property_age', 'community_pool', 'additional_features'
]


from routes.profile import profile_bp

def _query_monthly_contacts(user_id):
    """Retorna leads contactados en los últimos 30 días para un profesional."""
    conn = None
    try:
        conn = models.get_db_connection()
        rows = conn.execute('''
            SELECT l.id, l.type, l.property_type, l.zone, l.province,
                   l.budget, l.currency, l.timestamp,
                   lt.seen, lt.contacted, lt.seen_at, lt.contacted_at
            FROM leads l
            JOIN lead_tracking lt ON l.id = lt.lead_id
            WHERE lt.professional_id = ?
              AND lt.seen = 1
              AND lt.contacted = 1
              AND lt.contacted_at >= datetime('now', '-30 days')
            ORDER BY lt.contacted_at DESC
            LIMIT 1000
        ''', (user_id,)).fetchall()
        return [dict(r) for r in rows]
    finally:
        if conn:
            conn.close()



def _build_export_stats(leads):
    """Construye stats a partir de la lista de leads."""
    total = len(leads)
    budgets = []
    zones = {}
    currencies = {}
    types = {}
    budget_by_currency = {}
    for lead in leads:
        b = parse_budget(lead.get('budget'))
        budgets.append(b)
        z = lead.get('zone') or 'Sin zona'
        zones[z] = zones.get(z, 0) + 1
        c = lead.get('currency') or 'ARG'
        currencies[c] = currencies.get(c, 0) + 1
        t = lead.get('type') or 'Otro'
        types[t] = types.get(t, 0) + 1
        by_curr = budget_by_currency.setdefault(c, {'count': 0, 'total': 0.0})
        by_curr['count'] += 1
        by_curr['total'] += b
    avg_budget = round(sum(budgets) / total, 2) if total else 0
    total_budget = sum(budgets)
    sorted_zones = sorted(zones.items(), key=lambda x: x[1], reverse=True)
    return {
        'total': total,
        'avg_budget': avg_budget,
        'total_budget': total_budget,
        'zones': sorted_zones,
        'currencies': currencies,
        'types': types,
        'zone_count': len(zones),
        'budget_by_currency': budget_by_currency,
    }


@profile_bp.route('/api/profile/professional/export/xlsx')
@decorators.professional_required
def profile_export_history_xlsx():
    try:
        lang = get_language()
        user_id = session['user_id']
        leads = _query_monthly_contacts(user_id)
        stats = _build_export_stats(leads)

        wb = Workbook()
        data_font = Font(name='Manrope', size=10, color='000410')
        label_font = Font(name='Manrope', size=10, bold=True, color='735A3A')

        # ─── Sheet 1: Resumen ───
        ws_resumen = wb.active
        ws_resumen.title = t('profile.xlsx_summary', lang)
        resumen_data = [
            [t('profile.xlsx_total_leads_contacted', lang), stats['total']],
            [t('profile.xlsx_avg_budget', lang), f"${stats['avg_budget']:,.2f}"],
            [t('profile.xlsx_total_budget', lang), f"${stats['total_budget']:,.2f}"],
            [t('profile.xlsx_different_zones', lang), stats['zone_count']],
        ]
        for col, h in enumerate([t('profile.xlsx_metric', lang), t('profile.xlsx_value', lang)], 1):
            ws_resumen.cell(row=1, column=col, value=h)
        _style_header_row(ws_resumen, 2)
        ws_resumen.column_dimensions['A'].width = 32
        ws_resumen.column_dimensions['B'].width = 22
        for row_idx, row_data in enumerate(resumen_data, 2):
            ws_resumen.cell(row=row_idx, column=1, value=row_data[0]).font = label_font
            ws_resumen.cell(row=row_idx, column=2, value=row_data[1]).font = data_font
            ws_resumen.cell(row=row_idx, column=2).alignment = Alignment(horizontal='center')
            _apply_data_border(ws_resumen, row_idx, 2)

        note_row = len(resumen_data) + 3
        ws_resumen.cell(row=note_row, column=1,
            value=t('profile.xlsx_generated_by', lang)
        ).font = Font(name='Manrope', size=8, italic=True, color='A68A64')

        # ─── Sheet 2: Leads ───
        ws_leads = wb.create_sheet('Leads')
        headers = ['ID', t('profile.xlsx_type', lang), t('profile.xlsx_property', lang), t('profile.xlsx_zone', lang), t('profile.xlsx_province', lang), t('profile.xlsx_budget', lang),
                   t('profile.xlsx_currency', lang), t('profile.xlsx_created', lang), t('profile.xlsx_seen', lang), t('profile.xlsx_contacted', lang)]
        for col, h in enumerate(headers, 1):
            ws_leads.cell(row=1, column=col, value=h)
        _style_header_row(ws_leads, len(headers))
        col_widths = [8, 22, 16, 20, 18, 14, 10, 18, 18, 18]
        for i, w in enumerate(col_widths, 1):
            ws_leads.column_dimensions[get_column_letter(i)].width = w
        ws_leads.auto_filter.ref = f'A1:J{len(leads) + 1}'
        for row_idx, lead in enumerate(leads, 2):
            vals = [
                lead['id'], lead['type'], lead['property_type'], lead['zone'],
                lead.get('province', ''), lead['budget'], lead['currency'],
                lead.get('timestamp', ''), lead.get('seen_at', ''), lead.get('contacted_at', ''),
            ]
            for col, val in enumerate(vals, 1):
                cell = ws_leads.cell(row=row_idx, column=col, value=val)
                cell.font = data_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            _apply_data_border(ws_leads, row_idx, len(headers))

        # ─── Sheet 3: Por Zona ───
        ws_zones = wb.create_sheet('Por Zona')
        for col, h in enumerate([t('profile.xlsx_zone', lang), t('profile.xlsx_count', lang), t('profile.xlsx_pct_total', lang)], 1):
            ws_zones.cell(row=1, column=col, value=h)
        _style_header_row(ws_zones, 3)
        ws_zones.column_dimensions['A'].width = 30
        ws_zones.column_dimensions['B'].width = 14
        ws_zones.column_dimensions['C'].width = 14
        ws_zones.auto_filter.ref = f'A1:C{len(stats["zones"]) + 1}'
        for row_idx, (zone, count) in enumerate(stats['zones'], 2):
            pct = round(count / stats['total'] * 100, 1) if stats['total'] else 0
            ws_zones.cell(row=row_idx, column=1, value=zone).font = label_font
            ws_zones.cell(row=row_idx, column=2, value=count).font = data_font
            ws_zones.cell(row=row_idx, column=2).alignment = Alignment(horizontal='center')
            ws_zones.cell(row=row_idx, column=3, value=f'{pct}%').font = data_font
            ws_zones.cell(row=row_idx, column=3).alignment = Alignment(horizontal='center')
            _apply_data_border(ws_zones, row_idx, 3)

        # ─── Sheet 4: Inversiones ───
        ws_inv = wb.create_sheet('Inversiones')
        for col, h in enumerate([t('profile.xlsx_currency', lang), t('profile.xlsx_count', lang), t('profile.xlsx_total_budget_col', lang), t('profile.xlsx_avg_budget_col', lang)], 1):
            ws_inv.cell(row=1, column=col, value=h)
        _style_header_row(ws_inv, 4)
        ws_inv.column_dimensions['A'].width = 16
        ws_inv.column_dimensions['B'].width = 14
        ws_inv.column_dimensions['C'].width = 22
        ws_inv.column_dimensions['D'].width = 24
        ws_inv.auto_filter.ref = f'A1:D{len(stats["currencies"]) + 1}'
        row_idx = 2
        for currency, count in sorted(stats['currencies'].items()):
            bc = stats['budget_by_currency'].get(currency, {'count': 0, 'total': 0.0})
            total_c = bc['total']
            avg_c = round(total_c / bc['count'], 2) if bc['count'] else 0
            ws_inv.cell(row=row_idx, column=1, value=currency).font = label_font
            ws_inv.cell(row=row_idx, column=2, value=count).font = data_font
            ws_inv.cell(row=row_idx, column=2).alignment = Alignment(horizontal='center')
            ws_inv.cell(row=row_idx, column=3, value=f"${total_c:,.2f}").font = data_font
            ws_inv.cell(row=row_idx, column=3).alignment = Alignment(horizontal='center')
            ws_inv.cell(row=row_idx, column=4, value=f"${avg_c:,.2f}").font = data_font
            ws_inv.cell(row=row_idx, column=4).alignment = Alignment(horizontal='center')
            _apply_data_border(ws_inv, row_idx, 4)
            row_idx += 1

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        now_str = datetime.now().strftime('%Y%m%d_%H%M')
        return send_file(
            buffer,
            as_attachment=True,
            download_name=f'reporte_rendimiento_{now_str}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
    except Exception:
        logger.exception('Error exporting XLSX report')
        lang = get_language()
        return jsonify({'error': t('profile.xlsx_export_error', lang)}), 500


