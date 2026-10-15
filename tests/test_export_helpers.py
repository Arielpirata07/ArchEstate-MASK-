from openpyxl import Workbook

from services.export_helpers import _apply_data_border, _style_header_row, pdf_safe, pdf_val


class TestPdfSafe:
    def test_none_returns_empty(self):
        assert pdf_safe(None) == ''

    def test_ascii_passthrough(self):
        assert pdf_safe('Palermo 123') == 'Palermo 123'

    def test_accents_stripped(self):
        assert pdf_safe('Construcción') == 'Construccion'
        assert pdf_safe('Ñoño') == 'Nono'

    def test_symbols_replaced(self):
        assert pdf_safe('100€') == '100EUR'
        assert pdf_safe('a—b') == 'a-b'

    def test_numbers(self):
        assert pdf_safe(200000) == '200000'


class TestPdfVal:
    def test_default_on_empty(self):
        assert pdf_val(None) == '-'
        assert pdf_val('') == '-'
        assert pdf_val('', default='N/A') == 'N/A'

    def test_value_passthrough(self):
        assert pdf_val('Palermo') == 'Palermo'


class TestXlsxHelpers:
    def _sheet(self, cols=3):
        wb = Workbook()
        return wb.active

    def test_style_header_row(self):
        ws = self._sheet()
        _style_header_row(ws, 3)
        assert ws.cell(row=1, column=1).font.bold is True
        assert ws.cell(row=1, column=3).fill.start_color.rgb == '00000410' or \
            ws.cell(row=1, column=3).fill.start_color.rgb.endswith('000410')

    def test_apply_data_border(self):
        ws = self._sheet()
        ws.cell(row=2, column=1).value = 'x'
        _apply_data_border(ws, 2, 3)
        assert ws.cell(row=2, column=1).border.left.style == 'thin'
        assert ws.cell(row=2, column=3).border.right.style == 'thin'
