"""Convención de storage UTC para marcas de tiempo.

Todas las columnas de timestamp se guardan en UTC naive: SQLite escribe
CURRENT_TIMESTAMP en UTC, las ventanas SQL comparan contra datetime('now')
(UTC) y el display convierte a Argentina con convert_to_argentina_time.

Antes de la corrección, tracking.py escribía ART y models/pipeline.py hora
local del servidor: el mismo campo dependía de quién lo escribiera y del host.
En producción (UTC) el timeline y las ventanas SQL quedaban corridos 3 h, y el
feed del dashboard mostraba otra hora que el timeline para el mismo hecho.
"""
import io
import uuid
from datetime import datetime, timedelta

import openpyxl
import pytz
import pytest
from werkzeug.security import generate_password_hash

import models
from db_migrations import _backfill_first_contact
from services.sla import _created_local, _to_server_local, business_hours_between
from utils import utc_now

UTC_FMT = '%Y-%m-%d %H:%M:%S'
ART = pytz.timezone('America/Argentina/Buenos_Aires')


def _to_art(naive_utc, fmt='%d/%m/%Y %H:%M:%S'):
    """Espera de test independiente de utils: UTC naive -> hora Argentina."""
    return (datetime.strptime(naive_utc, UTC_FMT)
            .replace(tzinfo=pytz.UTC).astimezone(ART).strftime(fmt))


def _make_pro(db):
    unique = uuid.uuid4().hex[:8]
    name = f'Utc Pro {unique}'
    cur = db.execute(
        'INSERT INTO users (username, email, hash, role, phone, phone_e164, is_active)'
        ' VALUES (?, ?, ?, ?, ?, ?, 1)',
        (name, f'utc_{unique}@test.com', generate_password_hash('pass'),
         'professional', '+5491198765432', '+5491198765432')
    )
    user_id = cur.lastrowid
    db.execute(
        'INSERT INTO professionals (user_id, name, license, specialty, status)'
        ' VALUES (?, ?, ?, ?, ?)',
        (user_id, name, f'LIC-{unique}', 'general', 'approved')
    )
    db.commit()
    return user_id, name


def _make_lead(db, timestamp=None):
    unique = uuid.uuid4().hex[:8]
    cols = ('type, property_type, zone, province, budget, currency, phone,'
            ' email, user_id')
    vals = ['comprar', 'departamento', 'Palermo', 'Buenos Aires', '200000',
            'USD', '+5491112345678', f'utc_{unique}@test.com', 1]
    if timestamp is not None:
        cols += ', timestamp'
        vals.append(timestamp)
    cur = db.execute(
        f'INSERT INTO leads ({cols}) VALUES ({",".join(["?"] * len(vals))})',
        vals
    )
    db.commit()
    return cur.lastrowid


def _login(client, user_id, username, role='professional'):
    with client.session_transaction() as sess:
        sess['user_id'] = user_id
        sess['username'] = username
        sess['role'] = role


def _assert_utc_now(text, label):
    delta = abs((utc_now() - datetime.strptime(text, UTC_FMT)).total_seconds())
    # Tolerancia holgada por latencia de request, pero muy por debajo del
    # offset de 3 h que producía guardar hora local en un host ART.
    assert delta <= 120, f'{label} no está en UTC: {text} (delta {delta:.0f}s)'


class TestTrackingWritersStoreUtc:
    def test_toggle_seen_stores_utc(self, client, db, monkeypatch):
        # notify_lead_status_change intenta SMTP real y enlentece el request
        # (el assert mide la escritura, no el canal de email).
        import services.notifications as notifications
        monkeypatch.setattr(notifications, 'notify_lead_status_change',
                            lambda *a, **k: None)

        user_id, name = _make_pro(db)
        _login(client, user_id, name)
        lead_id = _make_lead(db)

        resp = client.post(f'/api/lead/{lead_id}/toggle-status',
                           json={'status': 'seen'})
        assert resp.status_code == 200

        row = db.execute(
            'SELECT seen_at FROM lead_tracking'
            ' WHERE professional_id = ? AND lead_id = ?',
            (user_id, lead_id)
        ).fetchone()
        _assert_utc_now(row['seen_at'], 'lead_tracking.seen_at')

    def test_move_lead_stage_stores_utc(self, db):
        user_id, _ = _make_pro(db)
        lead_id = _make_lead(db)

        moved = models.move_lead_stage(user_id, lead_id, 'contacted')
        assert moved is not None

        checks = {
            'lead_stage_history.changed_at':
                'SELECT changed_at AS ts FROM lead_stage_history WHERE lead_id = ?',
            'lead_stages.updated_at':
                'SELECT updated_at AS ts FROM lead_stages WHERE lead_id = ?',
            'lead_tracking.contacted_at':
                'SELECT contacted_at AS ts FROM lead_tracking WHERE lead_id = ?',
        }
        for label, sql in checks.items():
            row = db.execute(sql, (lead_id,)).fetchone()
            _assert_utc_now(row['ts'], label)


class TestDashboardFeedShowsArgentinaTime:
    def test_activity_at_is_argentina_wall_time(self, client, db):
        """El feed viaja como muro horario ART naive: dashTimeAgo() en JS lo
        parsea como hora local del navegador."""
        user_id, name = _make_pro(db)
        _login(client, user_id, name)
        lead_id = _make_lead(db)

        stage_ts = '2026-01-05 12:00:00'   # 09:00 ART
        contact_ts = '2026-01-05 15:30:00'  # 12:30 ART
        db.execute(
            'INSERT INTO lead_stage_history'
            ' (professional_id, lead_id, from_stage, to_stage, changed_at)'
            ' VALUES (?, ?, ?, ?, ?)',
            (user_id, lead_id, 'new', 'contacted', stage_ts)
        )
        db.execute(
            'INSERT INTO lead_tracking'
            ' (professional_id, lead_id, seen, contacted, seen_at, contacted_at)'
            ' VALUES (?, ?, 1, 1, ?, ?)',
            (user_id, lead_id, stage_ts, contact_ts)
        )
        db.commit()

        resp = client.get('/api/dashboard/activity?limit=20')
        assert resp.status_code == 200
        events = [e for e in resp.get_json()['activity']
                  if e['lead_id'] == lead_id]
        by_kind = {e['kind']: e for e in events}

        assert set(by_kind) == {'stage_move', 'contacted', 'seen'}
        assert by_kind['stage_move']['at'] == _to_art(stage_ts, UTC_FMT)
        assert by_kind['seen']['at'] == _to_art(stage_ts, UTC_FMT)
        assert by_kind['contacted']['at'] == _to_art(contact_ts, UTC_FMT)
        # El instante más reciente encabeza el feed.
        assert events[0]['kind'] == 'contacted'

    def test_move_and_contact_share_single_entry(self, client, db):
        """El dedup compara changed_at con contacted_at: la normalización a
        hora ART tiene que preservar la igualdad del valor crudo."""
        user_id, name = _make_pro(db)
        _login(client, user_id, name)
        lead_id = _make_lead(db)

        same_ts = '2026-01-05 12:00:00'
        db.execute(
            'INSERT INTO lead_stage_history'
            ' (professional_id, lead_id, from_stage, to_stage, changed_at)'
            ' VALUES (?, ?, ?, ?, ?)',
            (user_id, lead_id, 'new', 'contacted', same_ts)
        )
        db.execute(
            'INSERT INTO lead_tracking'
            ' (professional_id, lead_id, contacted, contacted_at)'
            ' VALUES (?, ?, 1, ?)',
            (user_id, lead_id, same_ts)
        )
        db.commit()

        resp = client.get('/api/dashboard/activity?limit=20')
        events = [e for e in resp.get_json()['activity']
                  if e['lead_id'] == lead_id]
        assert [e['kind'] for e in events] == ['stage_move']


class TestDaysInStageUsesUtcNow:
    def test_days_in_stage_against_utc_storage(self, client, db):
        user_id, name = _make_pro(db)
        _login(client, user_id, name)
        lead_id = _make_lead(db)

        # 5 días + 120 s en storage UTC. Restando hora local del servidor en
        # un host ART (−3 h) el resultado sería 4 días: el assert discrimina.
        updated_at = (utc_now() - timedelta(days=5, seconds=120)).strftime(UTC_FMT)
        db.execute(
            'INSERT INTO lead_stages (professional_id, lead_id, stage, updated_at)'
            ' VALUES (?, ?, ?, ?)',
            (user_id, lead_id, 'new', updated_at)
        )
        db.commit()

        resp = client.get('/api/pipeline/leads?my_leads=0')
        assert resp.status_code == 200
        leads = {l['id']: l for l in resp.get_json()['leads']}
        assert leads[lead_id]['days_in_stage'] == 5


class TestBackfillTreatsContactedAtAsUtc:
    def test_backfill_converts_contacted_at(self, db):
        # Jue 08/01/2026 09:00 ART = 12:00 UTC (creación) y lunes 12/01
        # 16:00 ART = 19:00 UTC (contacto): 9 h (jue) + 9 h (vie) + 7 h (lun)
        # = 25 h hábiles. Sin convertir contacted_at, 19:00 se leería como
        # hora local (ART) y daría 27 h.
        user_id, _ = _make_pro(db)
        lead_id = _make_lead(db, timestamp='2026-01-08 12:00:00')
        db.execute(
            'INSERT INTO lead_tracking'
            ' (professional_id, lead_id, seen, contacted, contacted_at)'
            ' VALUES (?, ?, 1, 1, ?)',
            (user_id, lead_id, '2026-01-12 19:00:00')
        )
        db.commit()

        before = {r['lead_id'] for r in
                  db.execute('SELECT lead_id FROM lead_first_contact').fetchall()}
        try:
            # La firma espera un cursor (así la invoca app_setup en la migración).
            _backfill_first_contact(db.cursor())
            row = db.execute(
                'SELECT professional_id, source, contact_at,'
                ' business_hours_elapsed FROM lead_first_contact'
                ' WHERE lead_id = ?',
                (lead_id,)
            ).fetchone()
            assert row is not None
            assert row['professional_id'] == user_id
            assert row['source'] == 'stage'
            # contacted_at se interpreta en UTC y se canoniza en tz servidor.
            assert row['contact_at'] == _to_server_local(
                datetime(2026, 1, 12, 19, 0)).strftime(UTC_FMT)
            expected = business_hours_between(
                _created_local('2026-01-08 12:00:00'),
                _to_server_local(datetime(2026, 1, 12, 19, 0)),
            )
            assert row['business_hours_elapsed'] == pytest.approx(expected)
        finally:
            # El backfill recorre TODOS los leads de la DB compartida de la
            # sesión: retirar las filas que esta prueba agregó para no ensuciar
            # los tests de KPI que corren después.
            now_ids = {r['lead_id'] for r in
                       db.execute('SELECT lead_id FROM lead_first_contact').fetchall()}
            for extra in now_ids - before:
                db.execute('DELETE FROM lead_first_contact WHERE lead_id = ?',
                           (extra,))
            db.commit()


class TestExportDisplaysArgentinaTimes:
    def test_xlsx_leads_sheet_converts_timestamps(self, client, db):
        user_id, name = _make_pro(db)
        _login(client, user_id, name)

        created_utc = (utc_now() - timedelta(days=2)).strftime(UTC_FMT)
        seen_utc = (utc_now() - timedelta(days=1)).strftime(UTC_FMT)
        contact_utc = (utc_now() - timedelta(hours=6)).strftime(UTC_FMT)
        lead_id = _make_lead(db, timestamp=created_utc)
        db.execute(
            'INSERT INTO lead_tracking'
            ' (professional_id, lead_id, seen, contacted, seen_at, contacted_at)'
            ' VALUES (?, ?, 1, 1, ?, ?)',
            (user_id, lead_id, seen_utc, contact_utc)
        )
        db.commit()

        resp = client.get('/api/profile/professional/export/xlsx')
        assert resp.status_code == 200

        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        ws = wb['Leads']
        # Hoja Leads: col 8 = creado, col 9 = visto, col 10 = contactado.
        assert ws.cell(row=2, column=8).value == _to_art(created_utc)
        assert ws.cell(row=2, column=9).value == _to_art(seen_utc)
        assert ws.cell(row=2, column=10).value == _to_art(contact_utc)
